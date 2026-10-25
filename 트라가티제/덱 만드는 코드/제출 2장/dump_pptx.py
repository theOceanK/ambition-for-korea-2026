import sys
from pptx import Presentation
from pptx.util import Emu

p = Presentation(sys.argv[1])
print("slide size", p.slide_width, p.slide_height)
for i, s in enumerate(p.slides, 1):
    print("=== slide", i, s.slide_layout.name)
    for sh in s.shapes:
        t = sh.shape_type
        txt = ""
        if sh.has_text_frame:
            paras = []
            for para in sh.text_frame.paragraphs:
                runs = [(r.text, r.font.size.pt if r.font.size else None, r.font.bold,
                         str(r.font.color.rgb) if r.font.color and r.font.color.type is not None and hasattr(r.font.color, 'rgb') else None,
                         r.font.name) for r in para.runs]
                paras.append(runs)
            txt = paras
        print(sh.shape_id, sh.name, t, round(sh.left/12700), round(sh.top/12700), round(sh.width/12700), round(sh.height/12700), txt)
