"""트라가티제 제출 2장 — Drive 템플릿(PPTX 내보내기)에 내용을 채운다.

사용:
  python build_pptx.py <템플릿.pptx> <출력.pptx> [--demo URL] [--video URL] [--shot 그림.png]

--demo 를 주면 데모 링크 자리에 넣고 QR 칸에 QR 그림을 만든다.
--shot 을 주면 데모 칸을 그 그림으로 채우고, 없으면 작동 흐름 그림을 그린다.
"""
import argparse
import copy
import io

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

FONT = "Noto Sans KR"

# ---- 내용 (원천 = 트라가티제_요약_2장.pdf) ---------------------------------
S1 = {
    42: ["TOPIC · Wildcard"],
    23: ["트라가티제",
         "Upload product info or a video. Global shoppers watch, call, and buy."],
    26: ["서울 성수·명동에 매장을 둔|2~9인 스몰 브랜드 대표",
         "네이버·쿠팡에서 자기 상품을 직접 판다"],
    29: ["성수 외국인 카드 소비 +103.7%인데,|귀국한 손님은 다시 살 길이 없다",
         "쇼핑몰은 한국 카드만, 대행은 월 80만 원+"],
    31: ["1.2조 원"],
    32: ["연 TAM 추정 · 한국 온라인 판매 사업자 70.2만 명|4년차부터 미국 틱톡샵 셀러로 확장"],
    35: ["지금 · 콘텐츠 대행", "월 80만 원+"],
    38: ["우리 제품 · 라이트", "월 19.9만 원"],
    39: ["출처: 서울시(2026.10, 성수 외국인 카드 소비 2026.1~8) · 국세청 공공데이터(2026.8)"
         " · 대행비·요금: 트라가티제 기획서"],
}
S2 = {
    49: ["영상은 다국어 숏폼으로, 전화는 AI가 받는다"],
    54: ["상품별 숏폼 · 다국어 자막 · AI 전화"],
    57: ["숏폼 · 전화 · 구매를 월정액 하나로"],
    65: ["팀원  김덕환 · 김성재 · 윤정민 · 김해청"],
}
FLOW = [
    ("① 콘텐츠", ["제품 정보나|영상 올리기", "상품별|다국어 숏폼", "보관함 ·|SNS 자동 게시"]),
    ("② 구매 연결", ["국내 손님|네이버 · 쿠팡", "해외 손님|틱톡샵 JP · 큐텐", "매장 손님|QR → 재구매"]),
    ("③ AI 전화", ["질문 응대", "사이즈 추천", "구매 링크|발송", "대표에게|한국어 요약"]),
]
FLOW_TITLE = "작동 흐름"
PRICE_LINE = "월정액 4.9만 · 19.9만 · 39.9만 원 · 포함 클릭 넘으면 1번 50원 · 판매 수수료 0"


# ---- 도우미 ---------------------------------------------------------------
def shape_by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def set_paras(shape, texts):
    """문단마다 첫 런의 서식을 살린 채 글만 바꾼다. 문단 수는 템플릿과 같아야 한다."""
    paras = shape.text_frame.paragraphs
    assert len(paras) == len(texts), (shape.shape_id, len(paras), len(texts))
    for p, t in zip(paras, texts):
        runs = p.runs
        parts = t.split("|")          # | = 줄바꿈(<a:br/>) — 한글 낱말이 줄 끝에서 쪼개지지 않게
        runs[0].text = parts[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        prev = runs[0]._r
        for part in parts[1:]:
            br = etree.Element(qn("a:br"))
            rpr = prev.find(qn("a:rPr"))
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            prev.addnext(br)
            nr = copy.deepcopy(runs[0]._r)
            nr.find(qn("a:t")).text = part
            br.addnext(nr)
            prev = nr


def set_font(run, size, bold, theme):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.theme_color = theme
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    """lines = [(글, 크기pt, 굵게, 테마색), ...] — 한 줄이 한 문단."""
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, (t, size, bold, theme) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = t
        set_font(r, size, bold, theme)
    return tb


def add_box(slide, x, y, w, h, fill, radius=0.12):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.theme_color = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def draw_flow(slide, bx, by, bw, bh):
    pad = 14
    add_text(slide, bx + pad, by + 10, 200, 18, [(FLOW_TITLE, 10.5, True, MSO_THEME_COLOR.DARK_2)],
             align=PP_ALIGN.LEFT)
    label_w = 62
    gap = 12           # 칸 사이(화살표 자리)
    row_h = 40
    top = by + 34
    for ri, (label, chips) in enumerate(FLOW):
        y = top + ri * (row_h + 9)
        add_text(slide, bx + pad, y, label_w, row_h, [(label, 10.5, True, MSO_THEME_COLOR.ACCENT_1)],
                 align=PP_ALIGN.LEFT)
        cx0 = bx + pad + label_w
        avail = bx + bw - pad - cx0
        n = len(chips)
        cw = (avail - gap * (n - 1)) / n
        for ci, t in enumerate(chips):
            cx = cx0 + ci * (cw + gap)
            last = ci == n - 1 and ri != 1   # 구매 연결 줄은 나란한 갈래라 강조 없음
            add_box(slide, cx, y, cw, row_h,
                    MSO_THEME_COLOR.DARK_2 if last else MSO_THEME_COLOR.LIGHT_1)
            fg = MSO_THEME_COLOR.LIGHT_1 if last else MSO_THEME_COLOR.DARK_1
            add_text(slide, cx + 4, y, cw - 8, row_h, [(part, 10.5, True, fg) for part in t.split("|")])
            if ci < n - 1 and ri != 1:
                add_text(slide, cx + cw, y, gap, row_h, [("→", 10.5, True, MSO_THEME_COLOR.ACCENT_1)])
    draw_price(slide, bx + pad, by + bh - 26, bw - 2 * pad)


def draw_price(slide, x, y, w):
    # 2장 노트 「어떻게 돈을 버는지」 — 캡처로 바꿔도 남긴다
    add_text(slide, x, y, w, 18, [(PRICE_LINE, 10.5, True, MSO_THEME_COLOR.DARK_2)], align=PP_ALIGN.LEFT)


def fit_picture(slide, path, x, y, w, h):
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(path, Pt(x), Pt(y), Pt(w), Pt(h))
    # 칸 비율에 맞춰 넘치는 쪽을 잘라 낸다(cover)
    box_r, img_r = w / h, iw / ih
    if img_r > box_r:
        c = (1 - box_r / img_r) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - img_r / box_r) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("out")
    ap.add_argument("--demo")
    ap.add_argument("--video")
    ap.add_argument("--shot")
    a = ap.parse_args()

    prs = Presentation(a.template)
    s1, s2 = prs.slides[0], prs.slides[1]

    for sid, texts in S1.items():
        set_paras(shape_by_id(s1, sid), texts)
    for sid, texts in S2.items():
        set_paras(shape_by_id(s2, sid), texts)

    # 데모 링크 상자
    card = shape_by_id(s2, 62)
    p = card.text_frame.paragraphs
    p[1].runs[0].text = a.demo.replace("https://", "") if a.demo else "[데모 링크]"
    # 긴 주소는 한 줄(약 170pt)에 들어가게 글자를 줄인다 — 하한 10.5pt(=28px)
    p[1].runs[0].font.size = Pt(max(10.5, min(14, 170 / (0.58 * len(p[1].runs[0].text)))))
    p[2].runs[1].text = a.video.replace("https://", "") if a.video else "[영상 링크]"
    if a.demo:
        import qrcode
        img = qrcode.make(a.demo, border=1)
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)
        qr = shape_by_id(s2, 63)
        s2.shapes.add_picture(buf, qr.left + Pt(3), qr.top + Pt(3), qr.width - Pt(6), qr.height - Pt(6))
        txt = shape_by_id(s2, 64)
        txt._element.getparent().remove(txt._element)

    # 데모 칸: 점선 지우고 그림 또는 작동 흐름
    box = shape_by_id(s2, 50)
    hint = shape_by_id(s2, 51)
    hint._element.getparent().remove(hint._element)
    bx, by, bw, bh = (v / 12700 for v in (box.left, box.top, box.width, box.height))
    if a.shot:
        box._element.getparent().remove(box._element)
        fit_picture(s2, a.shot, bx, by, bw, bh - 26)
        draw_price(s2, bx, by + bh - 18, bw)
    else:
        box.line.fill.background()
        draw_flow(s2, bx, by, bw, bh)

    prs.save(a.out)
    print("saved", a.out)


if __name__ == "__main__":
    main()
